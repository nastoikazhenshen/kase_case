import json
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from mvp.src.finance_validation import compute_base_scenario, load_scenario, validate


SOURCE_WORKBOOK = Path(
    "mvp/data/raw/archive/verification/inputs/"
    "BondScope_Dashboard_Corrected_M12_8_M24_10.xlsx"
)


class FinanceValidationTests(unittest.TestCase):
    def test_base_scenario_recomputes_all_reference_checkpoints(self):
        report = validate(load_scenario())
        self.assertTrue(report["valid"], report["mismatches"])
        self.assertEqual(8, report["computed"]["month_12_active_organizations"])
        self.assertEqual(10, report["computed"]["month_24_active_organizations"])
        self.assertEqual(96_600_000, report["computed"]["revenue_24m_kzt"])
        self.assertEqual(12_602_000, report["computed"]["cumulative_cashflow_m24_kzt"])
        self.assertEqual(20, report["computed"]["payback_month"])

    def test_result_depends_on_monthly_schedule_not_cached_checkpoint(self):
        scenario = load_scenario()
        scenario["monthly_new_organizations"] = [0] * 24
        result = compute_base_scenario(scenario)
        self.assertEqual(0, result["revenue_24m_kzt"])
        self.assertEqual(-71_228_000, result["cumulative_cashflow_m24_kzt"])
        self.assertIsNone(result["payback_month"])

    def test_source_declaration_identifies_assumption_not_actual_customers(self):
        path = Path("mvp/data/base_scenario.json")
        payload = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual("ASSUMPTION", payload["status"])
        self.assertEqual("paying organization", payload["unit"])
        self.assertIn("actual customers", payload["not_claimed"])

    @unittest.skipUnless(SOURCE_WORKBOOK.exists(), "local archived workbook is unavailable")
    def test_base_scenario_matches_archived_workbook_input_cells(self):
        """Bind the independent model inputs to the immutable workbook, not its totals."""
        scenario = load_scenario()
        workbook = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=False)
        inputs = workbook["04 Inputs"]
        sales_plan = workbook["15 Sales Plan"]
        budget = workbook["05 Build Budget"]

        self.assertEqual(inputs["C8"].value, scenario["monthly_subscription_kzt"])
        self.assertEqual(inputs["C15"].value, scenario["monthly_opex_kzt"])
        self.assertEqual(
            inputs["C16"].value + inputs["C17"].value,
            scenario["variable_cost_per_active_organization_kzt"],
        )
        self.assertEqual(inputs["C18"].value, scenario["connection_cost_per_new_organization_kzt"])
        self.assertEqual("=SUM(C22:H22)", budget["I22"].value)
        startup_from_detailed_inputs = 0
        for column in range(3, 9):
            letter = get_column_letter(column)
            direct_costs = sum(budget.cell(row=row, column=column).value for row in (10, 11, 12, 14, 15, 16, 17, 18))
            management = sum(budget.cell(row=row, column=column).value for row in (10, 11, 12)) * budget["C26"].value
            subtotal = direct_costs + management
            self.assertEqual(f"=SUM({letter}10:{letter}18)", budget.cell(row=20, column=column).value)
            self.assertEqual(f"={letter}20*$C$27", budget.cell(row=21, column=column).value)
            self.assertEqual(f"=SUM({letter}20:{letter}21)", budget.cell(row=22, column=column).value)
            startup_from_detailed_inputs += subtotal * (1 + budget["C27"].value)
        self.assertEqual(
            startup_from_detailed_inputs,
            scenario["startup_expenditure_kzt"],
        )
        self.assertEqual(
            [sales_plan.cell(row=row, column=4).value for row in range(6, 30)],
            scenario["monthly_new_organizations"],
        )
